"""Gera uma previa privada e conciliada; nao alimenta o site ou importadores."""

import argparse
import csv
import hashlib
import json
from collections import defaultdict
from decimal import Decimal, ROUND_HALF_UP
from pathlib import Path

from openpyxl import load_workbook


ROOT = Path(__file__).resolve().parents[1]
UNIDADES = {"37201": "SergipePrevidência", "37401": "FINANPREV", "37403": "FUNPREV"}
CENTAVO = Decimal("0.01")


def moeda(valor):
    return format(valor.quantize(CENTAVO, rounding=ROUND_HALF_UP), ".2f")


def soma(linhas, campo):
    return sum((linha[campo] for linha in linhas), Decimal(0))


def gravar_csv(destino, linhas):
    with destino.open("w", encoding="utf-8-sig", newline="") as arquivo:
        writer = csv.DictWriter(arquivo, fieldnames=list(linhas[0]), delimiter=";")
        writer.writeheader()
        writer.writerows(linhas)


def gerar(origem, destino):
    origem, destino = origem.resolve(), destino.resolve()
    if destino.is_relative_to(ROOT / "dist") or destino == origem.parent:
        raise ValueError("A previa deve ficar fora de dist e da pasta da origem.")
    hash_antes = hashlib.sha256(origem.read_bytes()).hexdigest()
    livro = load_workbook(origem, read_only=True, data_only=False)
    detalhes, mensal, anual, chaves = [], defaultdict(list), defaultdict(list), set()
    try:
        for aba in livro:
            if aba.max_column != 13:
                raise ValueError(f"{aba.title}: esperadas 13 colunas da exportacao.")
            for numero, valores in enumerate(aba.iter_rows(min_row=2, values_only=True), 2):
                if any(valor is None for valor in valores):
                    raise ValueError(f"{aba.title}:{numero}: celula vazia; revisar origem.")
                uo, ano, mes = str(valores[2]), int(valores[3]), int(valores[4])
                if uo not in UNIDADES or not 1 <= mes <= 12:
                    raise ValueError(f"{aba.title}:{numero}: UO ou mes nao reconhecido.")
                dinheiro = []
                for valor in valores[11:13]:
                    if isinstance(valor, bool) or not isinstance(valor, (int, float)):
                        raise ValueError(f"{aba.title}:{numero}: valor nao numerico/formula.")
                    decimal = Decimal(str(valor))
                    if not decimal.is_finite():
                        raise ValueError(f"{aba.title}:{numero}: valor nao finito.")
                    dinheiro.append(decimal)
                chave = (uo, ano, mes, str(valores[6]), str(valores[8]), str(valores[10]))
                if chave in chaves:
                    raise ValueError(f"{aba.title}:{numero}: classificacao duplicada; revisar.")
                chaves.add(chave)
                registro = dict(
                    aba=aba.title, linha=numero, codigo_uo=uo, unidade=UNIDADES[uo],
                    orgao_origem=valores[0], sigla_origem=valores[1], ano=ano, mes=mes,
                    categoria=valores[5], codigo_categoria=str(valores[6]),
                    origem=valores[7], codigo_origem=str(valores[8]),
                    especie=valores[9], codigo_especie=str(valores[10]),
                    previsto=dinheiro[0], arrecadado=dinheiro[1],
                )
                detalhes.append(registro)
                mensal[uo, ano, mes].append(registro)
                anual[uo, ano].append(registro)
    finally:
        livro.close()
    if not detalhes:
        raise ValueError("Planilha sem registros.")

    resumo_mensal = [dict(
        codigo_uo=uo, unidade=UNIDADES[uo], ano=ano, mes=mes,
        arrecadado_mensal=moeda(soma(linhas, "arrecadado")),
        arrecadado_decimal_origem=str(soma(linhas, "arrecadado")),
        quantidade_linhas=len(linhas),
        linhas_origem=" | ".join(f"{r['aba']}:{r['linha']}" for r in linhas),
    ) for (uo, ano, mes), linhas in sorted(mensal.items())]
    resumo_anual, conferencias = [], []
    for (uo, ano), linhas in sorted(anual.items()):
        janeiro = [r for r in linhas if r["mes"] == 1]
        if uo in ("37201", "37401"):
            if not janeiro or any(r["previsto"] != 0 for r in linhas if r["mes"] != 1):
                raise ValueError(f"{uo}/{ano}: previsao diverge da regra de janeiro.")
            previsao, regra = soma(janeiro, "previsto"), "Classificacoes de janeiro, uma vez"
        elif ano == 2026 and all(r["mes"] >= 6 and r["previsto"] == 0 for r in linhas):
            previsao, regra = Decimal(0), "Zero anual confirmado pela instituicao; inicio em junho"
        else:
            raise ValueError(f"{uo}/{ano}: regra de previsao ainda nao confirmada.")
        resumo_anual.append(dict(
            codigo_uo=uo, unidade=UNIDADES[uo], ano=ano, previsao_anual=moeda(previsao),
            fundamento=regra,
            linhas_janeiro=" | ".join(f"{r['aba']}:{r['linha']}" for r in janeiro),
        ))
        agrupado = sum((Decimal(r["arrecadado_decimal_origem"]) for r in resumo_mensal
                        if r["codigo_uo"] == uo and r["ano"] == ano), Decimal(0))
        diferenca = agrupado - soma(linhas, "arrecadado")
        diferenca_previsao = previsao - soma(linhas, "previsto")
        if diferenca != 0 or diferenca_previsao != 0:
            raise ValueError(f"{uo}/{ano}: falha na conciliacao.")
        conferencias.append(dict(
            codigo_uo=uo, ano=ano, linhas=len(linhas),
            meses=sorted({r["mes"] for r in linhas}),
            arrecadacao_linhas=str(soma(linhas, "arrecadado")),
            arrecadacao_agrupada=str(agrupado), diferenca=str(diferenca),
            previsao_linhas=str(soma(linhas, "previsto")), previsao_anual=str(previsao),
            diferenca_previsao=str(diferenca_previsao),
            diferenca_soma_meses_exibidos= str(sum((Decimal(r["arrecadado_mensal"])
                for r in resumo_mensal if r["codigo_uo"] == uo and r["ano"] == ano),
                Decimal(0)) - Decimal(moeda(agrupado))),
        ))
    hash_depois = hashlib.sha256(origem.read_bytes()).hexdigest()
    if hash_antes != hash_depois:
        raise ValueError("A origem mudou durante a leitura; previa nao gravada.")
    # Decimal(str(valor)) preserva os residuos de precisao lidos pelo openpyxl.
    # Arredondamento apenas nas colunas de exibicao; detalhes conservam decimais.
    exportados = [{k: str(v) if isinstance(v, Decimal) else v for k, v in r.items()}
                  for r in detalhes]
    negativos = [r for r in exportados if Decimal(r["arrecadado"]) < 0]
    relatorio = dict(
        origem=str(origem.relative_to(ROOT)) if origem.is_relative_to(ROOT) else str(origem),
        sha256_origem=hash_antes, origem_preservada=True,
        linhas=len(detalhes), grupos_mensais=len(mensal), grupos_anuais=len(anual),
        negativos=len(negativos),
        linhas_intraorcamentarias=sum("intraor" in r["categoria"].casefold() for r in detalhes),
        precisao="Decimal(str(valor openpyxl)); soma sem arredondar; exibicao em centavos ROUND_HALF_UP",
        consolidacao="Nenhum total entre UOs; nenhuma eliminacao de receitas intraorcamentarias",
        pendencias=["Conferencia contabil dos negativos", "Tratamento das intraorcamentarias",
                    "Fechamento de outubro de 2026", "Alteracoes posteriores da previsao"],
        conferencias_por_uo_ano=conferencias,
    )
    destino.mkdir(parents=True, exist_ok=True)
    gravar_csv(destino / "arrecadacao-mensal.csv", resumo_mensal)
    gravar_csv(destino / "previsao-anual.csv", resumo_anual)
    gravar_csv(destino / "classificacoes-origem.csv", exportados)
    if negativos:
        gravar_csv(destino / "arrecadacoes-negativas.csv", negativos)
    (destino / "conferencia.json").write_text(
        json.dumps(relatorio, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    print(f"Previa: {destino}; {len(detalhes)} linhas; {len(mensal)} grupos mensais; "
          f"{len(anual)} previsoes anuais; {len(negativos)} negativos. Conciliacao OK.")
    for r in negativos:
        print(f"Negativo {r['aba']}!M{r['linha']}: {r['codigo_uo']} "
              f"{r['mes']:02d}/{r['ano']}: {moeda(Decimal(r['arrecadado']))}")


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--origem", type=Path,
                        default=ROOT / "dados-entrada/receitas/receitas-teste.xlsx")
    parser.add_argument("--saida", type=Path, default=ROOT / "trabalho/receitas-previa")
    args = parser.parse_args()
    gerar(args.origem, args.saida)
