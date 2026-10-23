"""Alimenta somente Valores das Diárias a partir da tabela e observações fornecidas."""
from pathlib import Path
from collections import Counter
from decimal import Decimal
from html import escape
import argparse
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'dados-entrada/diarias/tabela-diarias.xlsx'
NOTES = ROOT / 'dados-entrada/diarias/observacoes-tabela-diarias.txt'
PAGE = ROOT / 'dist/Portal-da-Transparencia/Diarias/TabelaValores/index.html'
STATE = 'Para Dentro do Estado (Intraestadual)'
NATIONAL = 'Para Fora do Estado (Interestadual)'
INTERNATIONAL = 'Para Fora do País (Internacional)'


def load():
    book = load_workbook(SOURCE, read_only=True, data_only=False)
    try:
        sheet = book['base']
        assert tuple(next(sheet.values)) == ('ANO', 'TIPO', 'GRUPO', 'CARGO', 'VALOR')
        rows = list(enumerate(sheet.iter_rows(min_row=2, values_only=True), 2))
        for n, (_, kind, group, _, _) in rows:
            assert kind in (STATE, NATIONAL, INTERNATIONAL), f'Âmbito desconhecido: linha {n}'
            assert group in ('I', 'II', 'III', 'IV')
        return rows
    finally:
        book.close()


def decree_url():
    notes = NOTES.read_text(encoding='utf-8-sig')
    match = re.search(r'Decreto nº 239/2023[^\n]*?(https://[^\s]+)', notes)
    assert match, 'URL do Decreto nº 239/2023 não identificada nas observações'
    return match[1]


def money(number):
    value = Decimal(str(number))
    assert value == value.quantize(Decimal('0.01'))
    return format(value, ',.2f').replace(',', '_').replace('.', ',').replace('_', '.')


def value_html(kind, value):
    if kind == STATE:
        if value == '-':
            return '-'
        match = re.fullmatch(r'R\$ ([\d.,]+) com pernoite / R\$ ([\d.,]+) sem pernoite', value)
        assert match, f'Formato estadual não reconhecido: {value}'
        return f'Com pernoite: <span serigyItem="valorEstadual">R$ {escape(match[1])}</span><br>Sem pernoite: <span serigyItem="valorEstadual">R$ {escape(match[2])}</span>'
    if kind == NATIONAL:
        return f'<span serigyItem="valorNacional">R$ {money(value)}</span>'
    return f'<span serigyItem="valorInternacional">{money(value)}</span><br><small>€ na União Europeia; US$ nas demais localidades</small>'


def render(records):
    years = ''.join(f'<option value="{y}" serigyValue="{y}">{y}</option>' for y in sorted({r[0] for _, r in records}, reverse=True))
    rows = []
    for number, (year, kind, group, role, value) in records:
        rows.append(f'          <tr data-source-row="{number}" data-year="{year}"><td>{year}</td><td>{escape(kind)}</td><td>{escape(group)}</td><td>{escape(role)}</td><td>{value_html(kind, value)}</td></tr>')
    return f'''<main id="conteudo" class="container" serigySchema="https://serigy.tce.se.gov.br/comum/diarias.html">
    <h1>Valores das Diárias</h1>
    <p class="lead">Valores de referência por exercício, âmbito, grupo e cargo.</p>
    <p>Fundamento da tabela: <a id="decreto-valores-diarias" href="{escape(decree_url(), quote=True)}">Decreto nº 239/2023, de 17 de janeiro de 2023 (PDF)</a>, especialmente seu Anexo I.</p>
    <p id="moedas-valores-diarias">Valores dentro e fora do Estado em reais (R$). Fora do país: euros (€) para viagens à União Europeia e dólares (US$) para as demais localidades.</p>
    <form id="filtro-valores-diarias" class="filter-panel" role="search" aria-label="Consultar valores das diárias" data-table-filter data-target="tabela-valores-diarias" serigyItem="filtroPesquisa">
      <div class="field"><label for="ano-valores-diarias">Exercício</label><select id="ano-valores-diarias" name="ano" serigyItem="historicoExercicio"><option value="">Todos</option>{years}</select></div>
      <div class="filter-actions"><button id="pesquisar-valores-diarias" type="submit" serigyExecuteWaitSection="tabela-valores-diarias">Pesquisar</button><button type="reset" class="secondary" serigyClearForm>Limpar filtros</button></div>
    </form>
    <p class="result-count" data-count-for="tabela-valores-diarias" aria-live="polite">{len(records)} registro(s) encontrado(s).</p>
    <section id="resultados-valores-diarias" aria-labelledby="titulo-valores-diarias">
      <h2 id="titulo-valores-diarias">Valores de referência</h2>
      <div class="table-wrap"><table id="tabela-valores-diarias" serigyLoadSection="tabela-valores-diarias" aria-describedby="moedas-valores-diarias">
        <caption>Valores das diárias — grupos I a IV por exercício e âmbito</caption>
        <thead><tr><th scope="col">Exercício</th><th scope="col">Âmbito</th><th scope="col">Grupo</th><th scope="col">Cargo</th><th scope="col">Valor de referência</th></tr></thead>
        <tbody id="registros-valores-diarias">
{chr(10).join(rows)}
        </tbody>
      </table></div>
    </section>
    <section aria-labelledby="titulo-observacoes-diarias">
      <h2 id="titulo-observacoes-diarias">Observações</h2>
      <p>Dentro do Estado, os valores com e sem pernoite estão identificados separadamente.</p>
      <p>Fora do Estado: o Secretário Especial de Representação de Sergipe em Brasília e os respectivos servidores lotados nessa Secretaria não devem receber diárias pelo deslocamento dentro do Distrito Federal ou entre o Distrito Federal e o seu Estado de origem.</p>
    </section>
    <div class="update-info" serigyItem="dataUltimaAtualizacao">Informações atualizadas em <time datetime="2026-10-22">22/10/2026</time></div>
  </main>'''


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    records = load()
    page = PAGE.read_text(encoding='utf-8')
    page, count = re.subn(r'<main id="conteudo".*?</main>', lambda _: render(records), page, flags=re.S)
    assert count == 1
    page = page.replace('<body class="awaiting-information">', '<body>')
    if '/assets/js/tabelas.js' not in page:
        page = page.replace('</head>', '  <script defer src="/assets/js/tabelas.js"></script>\n</head>')
    if args.check:
        assert PAGE.read_bytes() == page.encode('utf-8'), 'HTML divergente da fonte'
    else:
        PAGE.write_bytes(page.encode('utf-8'))
    print({'registros': len(records), 'por_exercicio': dict(Counter(r[0] for _, r in records)), 'por_ambito': dict(Counter(r[1] for _, r in records))})


if __name__ == '__main__':
    main()
