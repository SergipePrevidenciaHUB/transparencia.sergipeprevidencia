"""Gera a relação de estágios e o CSV sem deduplicar períodos ou pessoas."""
from pathlib import Path
from datetime import datetime
from collections import Counter
import argparse
import csv
import io
import json
import re
from html import escape

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'dados-entrada/rh/estagiarios.xlsx'
PAGE = ROOT / 'dist/Portal-da-Transparencia/RecursosHumanos/Estagiarios/index.html'
CSV = ROOT / 'dist/dados/recursos-humanos/estagiarios.csv'
SCHEMA = 'https://serigy.tce.se.gov.br/comum/recursos-humanos.html'
MONTHS = ['Janeiro', 'Fevereiro', 'Março', 'Abril', 'Maio', 'Junho', 'Julho', 'Agosto', 'Setembro', 'Outubro', 'Novembro', 'Dezembro']
HEADERS = ['Ano de Exercício', 'Mês do Exercício', 'Nome', 'Setor/Local', 'Data de Contratação', 'Data de Término do Contrato (ou Desligamento)', 'Houve Aditivo?']


def load():
    book = load_workbook(SOURCE, read_only=True, data_only=False)
    try:
        sheet = book['Plan1']
        assert tuple(next(sheet.values)) == ('Ano', 'Mês', 'Nome', 'Setor/Local', 'Data de Contratação', 'Data de Término do Contrato (ou Desligamento)', 'Houve Aditivo?')
        records = []
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            year, month = str(int(row[0])), str(int(row[1]))
            assert 1 <= int(month) <= 12
            hired = datetime.strptime(row[4], '%d/%m/%Y')
            if row[5]:
                datetime.strptime(row[5], '%d/%m/%Y')
            values = [year, month] + [str(v) if v is not None else '' for v in row[2:]]
            records.append((number, str(hired.year), values))
        return records
    finally:
        book.close()


def csv_text(records):
    output = io.StringIO(newline='')
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_ALL, lineterminator='\r\n')
    writer.writerow(HEADERS)
    writer.writerows(values for _, _, values in records)
    return '\ufeff' + output.getvalue()


def render(records):
    def options(values):
        return ''.join(f'<option value="{v}" serigyValue="{v}">{v}</option>' for v in sorted(set(values), reverse=True))

    hiring = options(h for _, h, _ in records)
    years = options(v[0] for _, _, v in records)
    months = ''.join(f'<option value="{m}">{MONTHS[int(m)-1]}</option>' for m in sorted({v[1] for _, _, v in records}, key=int))
    rows = []
    for number, hired, values in records:
        y, m, name, sector, start, end, aditivo = values
        attr = escape(json.dumps(values, ensure_ascii=False), quote=True)
        cells = [f'<td>{y}</td>', f'<td>{MONTHS[int(m)-1]}</td>', f'<td serigyItem="nomeEstagiario">{escape(name)}</td>', f'<td>{escape(sector)}</td>', f'<td serigyItem="admissaoEstagiario">{escape(start)}</td>', f'<td serigyItem="terminoEstagiario">{escape(end)}</td>' if end else '<td>Não informado</td>', f'<td>{escape(aditivo) if aditivo else "Não informado"}</td>']
        rows.append(f'          <tr data-source-row="{number}" data-ano-contratacao="{hired}" data-ano-relacao="{y}" data-mes-relacao="{m}" data-export-values="{attr}">' + ''.join(cells) + '</tr>')
    return f'''<main id="conteudo" class="container">
    <h1>Estagiários</h1>
    <p class="lead">Relação de estagiários por período de referência.</p>
    <form id="filtro-estagiarios" class="filter-panel" serigySchema="{SCHEMA}" serigyItem="filtroPesquisa">
      <div class="field"><label for="estagiario-ano">Ano de contratação</label><select id="estagiario-ano" name="ano_contratacao"><option value="">Todos</option>{hiring}</select></div>
      <div class="field"><label for="estagiario-relacao-ano">Ano de Exercício</label><select id="estagiario-relacao-ano" name="ano_relacao" serigyItem="historicoExercicio"><option value="">Todos</option>{years}</select></div>
      <div class="field"><label for="estagiario-relacao-mes">Mês do Exercício</label><select id="estagiario-relacao-mes" name="mes_relacao"><option value="">Todos</option>{months}</select></div>
      <div class="filter-actions"><button type="submit" serigyExecuteWaitSection="estagiarioTabela">Pesquisar</button><button class="secondary" type="reset" serigyClearForm>Limpar filtros</button></div>
    </form>
    <div class="exports" serigySchema="{SCHEMA}"><strong>Baixar registros selecionados:</strong> <a id="exportar-estagiarios" href="/dados/recursos-humanos/estagiarios.csv" download serigyItem="salvarRelatorio" serigyValue="csv">CSV</a></div>
    <p class="result-count" data-count-for="estagiarioTabela" aria-live="polite">{len(records)} registro(s) encontrado(s).</p>
    <p id="estagiarios-vazio" hidden>Nenhum registro encontrado para os filtros selecionados.</p>
    <div class="table-wrap">
      <table id="estagiarioTabela" serigySchema="{SCHEMA}" serigyLoadSection="estagiarioTabela">
        <caption>Relação de estagiários — registros por ano e mês da relação</caption>
        <thead><tr><th scope="col">Ano de Exercício</th><th scope="col">Mês do Exercício</th><th scope="col">Nome completo</th><th scope="col">Setor/Local</th><th scope="col">Data de contratação</th><th scope="col">Data de término do contrato (ou desligamento)</th><th scope="col">Houve aditivo?</th></tr></thead>
        <tbody>
          <!-- DADOS_ESTAGIARIOS_INICIO -->
{chr(10).join(rows)}
          <!-- DADOS_ESTAGIARIOS_FIM -->
        </tbody>
      </table>
    </div>
    <div class="update-info" serigySchema="https://serigy.tce.se.gov.br/comum/recursos-humanos.html" serigyItem="dataUltimaAtualizacao">Informações atualizadas em <time datetime="2026-10-22">22/10/2026</time></div>
  </main>'''


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true', help='Confere os artefatos sem gravar.')
    args = parser.parse_args()
    records = load()
    page = PAGE.read_text(encoding='utf-8')
    page = re.sub(r'<main id="conteudo".*?</main>', lambda _: render(records), page, flags=re.S)
    page = re.sub(r'<meta name="description" content="[^"]*">', '<meta name="description" content="Relação de estagiários do SergipePrevidência por período, com consulta por ano de contratação e exportação CSV.">', page)
    page = page.replace('/assets/js/tabelas.js', '/assets/js/estagiarios.js')
    outputs = [(PAGE, page), (CSV, csv_text(records))]
    for path, content in outputs:
        if args.check:
            assert path.read_bytes() == content.encode('utf-8'), f'Divergência: {path}'
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content.encode('utf-8'))
    print(json.dumps({'registros': len(records), 'ano_relacao': dict(Counter(v[0] for _, _, v in records)), 'ano_contratacao': dict(Counter(h for _, h, _ in records)), 'terminos_vazios': sum(not v[5] for _, _, v in records)}, ensure_ascii=False))


if __name__ == '__main__':
    main()
