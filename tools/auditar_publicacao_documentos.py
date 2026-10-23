"""Audita destinos únicos e atualiza o controle geral sem alterar o portal.

Cache de evidências em trabalho/auditoria-documentos; --reverificar renova-o.
Campos manuais e abas preexistentes são preservados. GET limitado a 24 MiB,
quatro trabalhadores e no máximo duas conexões por host, sem HEAD/retries.
"""
import argparse, ast, concurrent.futures, csv, hashlib, io, json, re, socket, struct
import subprocess, tempfile, threading, time, zipfile
from collections import Counter, defaultdict
from datetime import datetime
from functools import partial
from html import unescape
from html.parser import HTMLParser
from http.server import SimpleHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path
from urllib.error import HTTPError
from urllib.parse import urljoin, urlsplit, urldefrag, unquote, quote
from urllib.request import Request, urlopen
from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.utils import get_column_letter
from pypdf import PdfReader

ROOT=Path(__file__).resolve().parents[1]
DIST=ROOT/'dist'
BASE='https://sergipeprevidencia.se.gov.br/'
WORK=ROOT/'trabalho/auditoria-documentos'
HEADERS=['Módulo','Página ou item','Descrição do documento','URL atual','Situação identificada',
         'Providência necessária','Status','Observações','Nome exato esperado do arquivo',
         'URL final após redirecionamento','Evidência da verificação','Data da conferência']
URL=re.compile(r'https?://[^\s<>"\x27`\\]+')
DOC=re.compile(r'\.(pdf|xlsx?|csv|json|docx?|odt|ods|zip)(?:$|[?#])',re.I)
now=datetime.now().astimezone().isoformat(timespec='seconds')
locks=defaultdict(lambda:threading.Semaphore(2))

class CDP:
    """Transporte WebSocket mínimo, com limite de quadro e fragmentação."""
    def __init__(self,url):
        u=urlsplit(url);self.s=socket.create_connection((u.hostname,u.port),timeout=30);self.id=0
        self.s.sendall(f'GET {u.path} HTTP/1.1\r\nHost: {u.netloc}\r\nUpgrade: websocket\r\nConnection: Upgrade\r\nSec-WebSocket-Key: dGhlIHNhbXBsZSBub25jZQ==\r\nSec-WebSocket-Version: 13\r\n\r\n'.encode())
        header=b''
        while not header.endswith(b'\r\n\r\n'):header+=self.read(1)
        assert header.startswith(b'HTTP/1.1 101 ')
    def read(self,n):
        if n>32*1024*1024:raise ValueError('Quadro CDP acima do limite de segurança')
        out=bytearray()
        while len(out)<n:
            part=self.s.recv(min(65536,n-len(out)))
            if not part:raise ConnectionError('CDP encerrado')
            out.extend(part)
        return bytes(out)
    def send(self,body,opcode=1):
        mask=b'abcd';n=len(body);head=bytes([128|opcode])
        head+=bytes([128|n]) if n<126 else bytes([254])+struct.pack('!H',n) if n<65536 else bytes([255])+struct.pack('!Q',n)
        self.s.sendall(head+mask+bytes(v^mask[i%4] for i,v in enumerate(body)))
    def call(self,method,params=None):
        self.id+=1;self.send(json.dumps(dict(id=self.id,method=method,params=params or {})).encode());parts=bytearray()
        while True:
            h=self.read(2);n=h[1]&127
            if n==126:n=struct.unpack('!H',self.read(2))[0]
            elif n==127:n=struct.unpack('!Q',self.read(8))[0]
            mask=self.read(4) if h[1]&128 else None;body=self.read(n)
            if mask:body=bytes(v^mask[i%4] for i,v in enumerate(body))
            opcode=h[0]&15
            if opcode==9:self.send(body,10);continue
            if opcode==10:continue
            if opcode==8:raise ConnectionError('CDP fechado')
            parts.extend(body)
            if not h[0]&128:continue
            msg=json.loads(parts);parts.clear()
            if msg.get('id')==self.id:
                if 'error' in msg:raise RuntimeError(str(msg['error']))
                return msg.get('result',{})
    def js(self,code):
        r=self.call('Runtime.evaluate',dict(expression=code,returnByValue=True,awaitPromise=True))
        if 'exceptionDetails' in r:raise RuntimeError(str(r['exceptionDetails']))
        return r.get('result',{}).get('value')

class Links(HTMLParser):
    def __init__(self):
        super().__init__();self.links=[];self.ids=set();self.active=None;self.row=[];self.rowlinks=[]
    def handle_starttag(self,t,a):
        a=dict(a)
        if a.get('id'):self.ids.add(a['id'])
        if t=='tr':self.row=[];self.rowlinks=[]
        for k in ['href','src','serigylink']:
            if a.get(k):
                item={'url':a[k],'context':a.get('aria-label') or a.get('alt') or a.get('title') or '', 'attribute':k}
                self.links.append(item);self.rowlinks.append(item)
                if t=='a':self.active=item
    def handle_data(self,s):
        self.row.append(s)
        if self.active:self.active['context']+=' '+s
    def handle_endtag(self,t):
        if t=='a':self.active=None
        if t=='tr':
            text=' '.join(' '.join(self.row).split())[:1800]
            for item in self.rowlinks:item['context']=text
            self.rowlinks=[];self.row=[]

def text_url(value):
    return unescape(value).split('](')[0].rstrip('.,;`)')

def module(page):
    parts=page.replace('\\','/').split('/')
    try:
        name=parts[parts.index('Portal-da-Transparencia')+1]
        return {'index.html':'Entrada','Convenios':'Convênios','Licitacoes':'Licitações','RecursosHumanos':'RH','PlanejamentoPrestacaoContas':'Planejamento','LGPDeGovernoDigital':'LGPD','Diarias':'Diárias'}.get(name,name)
    except (ValueError,IndexError):
        for token,label in [('convenio','Convênios'),('contrato','Contratos'),('ordem-cronologica','Contratos'),('planejamento','Planejamento'),('licitac','Licitações'),('portaria','Institucional'),('atos-normativos','Institucional'),('resolu','Institucional'),('institucional','Institucional'),('sic','SIC'),('ouvidoria','Ouvidoria'),('lgpd','LGPD'),('pesquisa','LGPD'),('diaria','Diárias'),('recursos-humanos','RH')]:
            if token in page.lower():return label
        return 'Controle interno'

def collect():
    occurrences=[]; pages=sorted((DIST/'Portal-da-Transparencia').rglob('*.html'))
    def add(url,page,context,kind):
        if not url or url.startswith(('mailto:','tel:','javascript:','data:','blob:','https://wa.me/')):return
        raw=text_url(url)
        resolved=urljoin(BASE+page,raw) if kind!='controle' and not raw.startswith('http') else raw
        if not resolved.startswith(('http://','https://')):return
        occurrences.append({'url':resolved,'page':page,'context':context.strip(),'kind':kind,'module':module(page)})
    scripts=set()
    for p in pages:
        page=p.relative_to(DIST).as_posix();parser=Links();parser.feed(p.read_text(encoding='utf-8-sig'))
        for item in parser.links:
            add(item['url'],page,item['context'],item['attribute'])
            if item['url'].endswith('.js') and item['url'].startswith('/'):
                scripts.add(DIST/item['url'].lstrip('/'))
    for p in sorted(scripts):
        if not p.exists():continue
        text=p.read_text(encoding='utf-8-sig')
        for match in URL.finditer(text):add(match.group(),p.relative_to(DIST).as_posix(),text[max(0,match.start()-120):match.end()+120],'javascript')
        for match in re.finditer(r'["\x27](/(?:Portal-da-Transparencia|dados|assets)/[^"\x27\n]+)["\x27]',text):
            add(match.group(1),p.relative_to(DIST).as_posix(),'Literal de caminho no JavaScript','javascript')
    # Evidência integral da linha/objeto fica preservada; nenhum dado de controle vai a dist.
    controls=[]
    for folder in ['controle','dados-entrada']:
        for p in (ROOT/folder).rglob('*'):
            if not p.is_file() or p.suffix.lower() not in ['.md','.txt','.json','.csv']:continue
            if folder=='dados-entrada' and not any(w in p.name.lower() for w in ['publica','document','link','preparad','revisad','consolidado','complement']):continue
            text=p.read_text(encoding='utf-8-sig',errors='replace')
            if p.suffix=='.json':
                try:
                    def walk(value,path=''):
                        if isinstance(value,dict):
                            for k,v in value.items():
                                if isinstance(v,str) and URL.search(v):
                                    context=json.dumps({x:y for x,y in value.items() if not isinstance(y,(dict,list))},ensure_ascii=False)
                                    for m in URL.finditer(v):add(m.group(),p.relative_to(ROOT).as_posix(),context[:12000],'controle')
                                elif isinstance(v,(dict,list)):walk(v,path+'/'+str(k))
                        elif isinstance(value,list):
                            for v in value:walk(v,path)
                    walk(json.loads(text));continue
                except (ValueError,TypeError):pass
            for n,line in enumerate(text.splitlines(),1):
                for m in URL.finditer(line):add(m.group(),p.relative_to(ROOT).as_posix()+':'+str(n),line[:12000],'controle')
                if folder=='controle' and not URL.search(line) and re.search(r'\.pdf|edital|regimento|balancete|acórdão',line,re.I) and re.search(r'ausent|falt|pend|obter|aguarda|não.*(?:fornecid|disponib)',line,re.I):
                    occurrences.append(dict(url='',page=p.relative_to(ROOT).as_posix()+':'+str(n),context=line[:12000],kind='documento sem URL',module=module(str(p))))
            controls.append(p)
    return pages,occurrences

def rendered(pages):
    """Executa os scripts atuais e extrai href/src/serigyLink do DOM inicial."""
    class Handler(SimpleHTTPRequestHandler):
        def log_message(self,*a):pass
        def do_GET(self):
            try:super().do_GET()
            except (BrokenPipeError,ConnectionAbortedError,ConnectionResetError):pass
    server=ThreadingHTTPServer(('127.0.0.1',0),partial(Handler,directory=str(DIST)))
    threading.Thread(target=server.serve_forever,daemon=True).start()
    probe=socket.socket();probe.bind(('127.0.0.1',0));port=probe.getsockname()[1];probe.close()
    profile=tempfile.mkdtemp(prefix='chrome-',dir=WORK)
    proc=subprocess.Popen([r'C:\Program Files\Google\Chrome\Application\chrome.exe','--headless','--disable-gpu','--no-sandbox','--no-first-run','--remote-allow-origins=*',f'--remote-debugging-port={port}','--user-data-dir='+profile,'about:blank'],stdout=subprocess.DEVNULL,stderr=subprocess.DEVNULL)
    out=[];done=0
    try:
        for _ in range(100):
            try:
                with urlopen(f'http://127.0.0.1:{port}/json',timeout=1) as r:targets=json.load(r)
                break
            except Exception:time.sleep(.2)
        c=CDP(next(t['webSocketDebuggerUrl'] for t in targets if t['type']=='page'));c.call('Page.enable')
        for p in pages:
            page=p.relative_to(DIST).as_posix();route='/'+page
            c.call('Page.navigate',{'url':f'http://127.0.0.1:{server.server_port}'+route})
            for _ in range(100):
                if c.js('location.pathname')==route and c.js('document.readyState')=='complete':break
                time.sleep(.05)
            else:continue
            items=c.js('''Array.from(document.querySelectorAll('[href],[src],[serigyLink]')).flatMap(e=>['href','src','serigyLink'].filter(a=>e.hasAttribute(a)).map(a=>({url:e.getAttribute(a),context:(e.closest('tr')||e).textContent.trim().slice(0,1800),kind:'DOM '+a})))''')
            for item in items:
                raw=item['url']
                if raw.startswith(('mailto:','tel:','data:','javascript:','blob:')):continue
                u=urljoin(BASE+page,raw)
                if u.startswith('http'):out.append(dict(item,url=u,page=page,module=module(page)))
            exports=c.js('''(async()=>{const output=[];for(const b of document.querySelectorAll('button[serigyItem="salvarRelatorio"]')){let blob;const create=URL.createObjectURL,click=HTMLAnchorElement.prototype.click;URL.createObjectURL=x=>{blob=x;return 'blob:auditoria-local'};HTMLAnchorElement.prototype.click=function(){};try{if(!b.disabled)b.click();await new Promise(r=>setTimeout(r,30));const text=blob&&blob.size<=25165824?await blob.text():'';let jsonValid=false;try{JSON.parse(text);jsonValid=true}catch(e){};output.push({label:b.textContent.trim(),disabled:b.disabled,type:blob?blob.type:'',bytes:blob?blob.size:0,text:text.slice(0,1000),lines:text.split(String.fromCharCode(10)).length,jsonValid});}finally{URL.createObjectURL=create;HTMLAnchorElement.prototype.click=click;}}return output})()''')
            for export in exports:
                if export['disabled']:state='Inconclusivo — exportação desativada no recorte inicial'
                elif not export['bytes']:state='Inconclusivo — controle sem arquivo Blob capturado'
                elif 'csv' in export['type']:
                    state='Exportação CSV gerada e lida: '+str(export['lines'])+' linhas físicas no recorte inicial'
                elif 'json' in export['type']:
                    state='Exportação JSON gerada e parseada' if export['jsonValid'] else 'Inconclusivo — JSON não interpretado'
                else:state='Inconclusivo — formato de exportação a conferir'
                out.append(dict(url='',page=page,context=export['label']+'; '+state+'; '+str(export['bytes'])+' bytes; SHA256 do início '+hashlib.sha256(export['text'].encode()).hexdigest(),kind='exportação gerada',module=module(page)))
            done+=1
        return out,f'DOM inicial executado no Chrome: {done}/{len(pages)} HTMLs; botões de exportação acionados com captura local de Blob, sem publicação ou submissão de formulários.'
    except Exception as e:return out,f'DOM Chrome inconclusivo após {done}/{len(pages)} páginas: {type(e).__name__}: {e}'
    finally:proc.terminate();server.shutdown();server.server_close()

def inspect_body(body,url,mime,truncated=False):
    suffix=Path(urlsplit(url).path).suffix.lower();evidence=[];title='';ids=[];state='Disponível';action='';status='Concluído'
    if truncated:return dict(state='Inconclusivo — limite de 24 MiB',action='Conferir arquivo grande manualmente.',status='Pendente',evidence='Corpo excede o limite; não tratado como ausente.',title='',ids=[])
    if not body.startswith(b'%PDF-') and re.search(rb'cf-chl-|<title[^>]*>\s*(?:sign in|login|entrar)|accounts\.google\.com/ServiceLogin',body[:200000],re.I):
        return dict(state='Inconclusivo — autenticação ou bloqueio',action='Conferir acesso público sem login/CAPTCHA.',status='Pendente',evidence='Resposta contém página de autenticação ou desafio; não comprova ausência do documento.',title='',ids=[])
    if body.startswith(b'%PDF-'):
        try:
            reader=PdfReader(io.BytesIO(body));title=str((reader.metadata or {}).get('/Title',''))
            texts=[p.extract_text() or '' for p in list(reader.pages)[:5]]
            title+=' '+' '.join(texts)[:20000];evidence.append(f'Assinatura PDF; {len(reader.pages)} páginas; texto inicial extraído.')
            if not ''.join(texts).strip():state='Inconclusivo — identificação em PDF sem texto';action='Conferir identificação visualmente.';status='Pendente'
        except Exception as e:state='Inconclusivo — PDF não interpretado';action='Abrir e conferir documento.';status='Pendente';evidence.append(type(e).__name__)
    elif suffix=='.pdf':state='Falha — destino não é PDF';action='Publicar o PDF correto no destino exato.';status='Pendente';evidence.append('Sem assinatura %PDF-; Content-Type: '+mime)
    elif suffix in ['.xlsx','.docx','.ods','.odt','.zip']:
        if not zipfile.is_zipfile(io.BytesIO(body)):state='Falha — formato diferente do anunciado';action='Publicar arquivo no formato indicado.';status='Pendente'
        else:
            with zipfile.ZipFile(io.BytesIO(body)) as z:
                names=z.namelist();expected={'.xlsx':'xl/workbook.xml','.docx':'word/document.xml'}
                if suffix in expected and expected[suffix] not in names:state='Falha — contêiner de outro formato';action='Publicar o arquivo correto.';status='Pendente'
            evidence.append('Assinatura ZIP e estrutura interna verificadas.')
    elif suffix in ['.xls','.doc']:
        if not body.startswith(bytes.fromhex('D0CF11E0A1B11AE1')):state='Inconclusivo — XLS sem contêiner OLE';action='Conferir formato real da planilha.';status='Pendente'
        evidence.append('Formato binário conferido; conteúdo não interpretado.')
    elif suffix=='.json':
        try:json.loads(body);evidence.append('JSON parseado.')
        except Exception:state='Falha — JSON inválido';action='Corrigir arquivo.';status='Pendente'
    elif suffix=='.csv':
        text=body.decode('utf-8-sig',errors='replace')
        if '<html' in text[:3000].lower():state='Falha — HTML no lugar de CSV';action='Publicar CSV correto.';status='Pendente'
        else:evidence.append(f'CSV lido: {len(list(csv.reader(io.StringIO(text))))} linhas; sem validar semântica financeira.')
    elif 'html' in mime or b'<html' in body[:2000].lower():
        text=body.decode('utf-8',errors='replace');p=Links();p.feed(text);ids=sorted(p.ids)
        m=re.search(r'<title[^>]*>(.*?)</title>',text,re.I|re.S);title=unescape(re.sub('<[^>]*>',' ',m[1])) if m else ''
        visible=re.sub(r'<(script|style)\b.*?</\1>','',text,flags=re.S|re.I);visible=re.sub('<[^>]*>',' ',visible)
        if re.search(r'captcha|cf-chl-|verify you are human|access denied',text,re.I):state='Inconclusivo — bloqueio/CAPTCHA';action='Verificar acesso manualmente.';status='Pendente'
        elif re.search(r'página não encontrada|page not found|documento não encontrado|arquivo não encontrado',title+' '+visible[:4000],re.I):state='Falha — mensagem de não encontrado em HTML';action='Corrigir destino ou publicar documento.';status='Pendente'
        elif len(visible.strip())<25:state='Inconclusivo — página vazia';action='Conferir carregamento e autenticação.';status='Pendente'
        evidence.append('HTML interpretado; título: '+title[:200])
    elif not body:state='Falha — resposta vazia';action='Disponibilizar conteúdo.';status='Pendente'
    else:evidence.append('Conteúdo recebido; tipo '+mime+'; identificação documental não interpretada.')
    return dict(state=state,action=action,status=status,evidence=' '.join(evidence),title=title,ids=ids)

def check(url,refresh):
    key=hashlib.sha256(url.encode()).hexdigest();cache=WORK/(key+'.json')
    if cache.exists() and not refresh:
        cached=json.loads(cache.read_text(encoding='utf-8'))
        if not (url==BASE and cached.get('evidence','').startswith('Pacote local')):return cached
    start=time.monotonic();result={'url':url,'final':url,'date':now}
    parsed=urlsplit(url);local=parsed.netloc=='sergipeprevidencia.se.gov.br' and parsed.path.startswith(('/Portal-da-Transparencia/','/assets/','/dados/','/mapa-do-site/'))
    try:
        if local:
            p=DIST/unquote(parsed.path).lstrip('/')
            if p.is_dir():p=p/'index.html'
            if not p.is_file():raise FileNotFoundError('Caminho ausente no pacote local: '+str(p.relative_to(ROOT)))
            body=p.read_bytes();mime='text/html' if p.suffix=='.html' else ''
            result.update(inspect_body(body,url,mime));result['evidence']='Pacote local atual (não comprova publicação remota). '+result['evidence']
        else:
            with locks[parsed.netloc]:
                time.sleep(.15)
                request_url=quote(url,safe=":/?&=%;,+@!$'()*[]")
                with urlopen(Request(request_url,headers={'User-Agent':'Portal-SergipePrevidencia-LinkAudit/1.0','Accept-Encoding':'identity'}),timeout=18) as response:
                    body=response.read(24*1024*1024+1);result['final']=response.url
                    mime=response.headers.get('Content-Type','');result.update(inspect_body(body,url,mime,len(body)>24*1024*1024))
                    result['evidence']=f'GET HTTP {response.status}; {mime}; {len(body)} bytes; SHA256 '+hashlib.sha256(body).hexdigest()+'. '+result['evidence']
                    if urlsplit(response.url).path.lower().startswith(('/login','/signin')):result.update(state='Inconclusivo — autenticação',action='Conferir acesso público sem login.',status='Pendente')
                    if response.url!=url and DOC.search(url) and not DOC.search(response.url) and not body.startswith(b'%PDF-'):result.update(state='Falha — redirecionamento para página sem documento',action='Corrigir destino documental.',status='Pendente')
    except HTTPError as e:
        result.update(final=e.url,state=('Ausência comprovada — HTTP '+str(e.code)) if e.code in [404,410] else 'Inconclusivo — HTTP '+str(e.code),action='Publicar no destino exato ou corrigir vínculo.' if e.code in [404,410] else 'Reverificar acesso/servidor sem presumir ausência.',status='Pendente',evidence=str(e),title='',ids=[])
    except FileNotFoundError as e:result.update(state='Ausência comprovada — arquivo local',action='Providenciar recurso no pacote.',status='Pendente',evidence=str(e),title='',ids=[])
    except Exception as e:result.update(state='Inconclusivo — acesso',action='Reverificar timeout, rede, bloqueio ou autenticação.',status='Pendente',evidence=type(e).__name__+': '+str(e),title='',ids=[])
    result['seconds']=round(time.monotonic()-start,2);cache.write_text(json.dumps(result,ensure_ascii=False),encoding='utf-8');return result

def write_book(path,occurrences,results,pages,browser):
    wb=load_workbook(path) if path.exists() else Workbook();existing={}
    if 'Pendências' in wb:
        ws=wb['Pendências'];heads=[c.value for c in ws[1]]
        for row in ws.iter_rows(min_row=2,values_only=True):
            d=dict(zip(heads,row));key=str(d.get('URL atual') or '')
            if not key:key='__semurl__'+hashlib.sha256(str((d.get('Página ou item'),d.get('Descrição do documento'))).encode()).hexdigest()
            existing[key]=d
    for name in ['Pendências','Ocorrências','Resumo']:
        if name in wb:del wb[name]
    if 'Sheet' in wb and wb['Sheet'].max_row==1:del wb['Sheet']
    ws=wb.create_sheet('Pendências');extra=[k for v in existing.values() for k in v if k not in HEADERS and k!='Histórico de conferências'];extra=list(dict.fromkeys(extra));heads=HEADERS+extra+['Histórico de conferências'];ws.append(heads)
    grouped=defaultdict(list)
    for o in occurrences:grouped[o['url']].append(o)
    counts=Counter();inconclusive=0
    for url in sorted(set(grouped)|set(existing)):
        items=grouped[url];old=existing.get(url,{})
        if url.startswith('__semurl__') and any(o['url']=='' and o['page']==old.get('Página ou item') and o['context'][:1000]==old.get('Descrição do documento') for o in occurrences):continue
        if not url:
            for item in items:
                row={'Módulo':item['module'],'Página ou item':item['page'],'Descrição do documento':item['context'][:1000],'Situação identificada':'Pendência documental interna — URL não fornecida','Providência necessária':'Identificar o documento e confirmar a providência indicada no controle original.','Status':'Pendente','Observações':item['context'],'Evidência da verificação':item['page'],'Data da conferência':now}
                if item['kind']=='exportação gerada':
                    okay='Inconclusivo' not in item['context']
                    row.update({'Situação identificada':'Exportação local verificada' if okay else 'Inconclusivo — exportação','Providência necessária':'' if okay else 'Conferir controle no recorte disponível.','Status':'Concluído' if okay else 'Pendente','Evidência da verificação':item['context']})
                    if not okay:inconclusive+=1
                key='__semurl__'+hashlib.sha256(str((row['Página ou item'],row['Descrição do documento'])).encode()).hexdigest()
                saved=existing.get(key,{})
                for field in ['Providência necessária','Status','Observações']+extra:
                    if saved.get(field):row[field]=saved[field]
                ws.append([row.get(k) for k in heads])
                if row['Status']!='Concluído':counts[item['module']]+=1
            continue
        base,fragment=urldefrag(url);r=results.get(base)
        if not r:
            ws.append([old.get(k) for k in heads]);continue
        r=dict(r);context=' | '.join(dict.fromkeys(o['context'] for o in items if o['context']))
        public=[o for o in items if o['kind']!='controle'];item=(public or items)[0]
        if fragment and r['status']=='Concluído' and not DOC.search(base) and unquote(fragment) not in r.get('ids',[]):r.update(state='Falha — âncora ausente',action='Corrigir âncora ou destino.',status='Pendente')
        # Identificação conservadora: só confirmar coincidência; ausência de texto é inconclusiva.
        ident=re.search(r'(?:portaria|resolução|lei|decreto|convênio|acordo)\s*(?:n[ºo°.]*)?\s*(\d{1,5})\s*[/.-]\s*(20\d{2}|19\d{2})',context,re.I)
        if DOC.search(base) and r['status']=='Concluído':
            if ident and r.get('title'):
                n,y=ident.groups();plain=re.sub(r'\s+',' ',r['title'])
                if not re.search(r'\b0?'+re.escape(str(int(n)))+r'\b',plain) or y not in plain:r.update(state='Inconclusivo — identidade documental',action='Conferir número/exercício no documento recebido; não corrigir pelo nome do arquivo.',status='Pendente')
            elif urlsplit(base).path.lower().endswith('.pdf') and not r.get('title'):r.update(state='Inconclusivo — identidade documental',action='Conferir identificação do PDF visualmente.',status='Pendente')
            if public and urlsplit(base).path.lower().endswith('.pdf') and r.get('title'):
                announced=' '.join(o['context'] for o in public)
                years=set(re.findall(r'\b(?:19|20)\d{2}\b',announced))
                if len(years)==1 and next(iter(years)) not in r['title']:
                    r.update(state='Inconclusivo — exercício não confirmado no documento',action='Conferir exercício anunciado com o PDF recebido.',status='Pendente')
                months=['janeiro','fevereiro','março','abril','maio','junho','julho','agosto','setembro','outubro','novembro','dezembro']
                stated=[m for m in months if re.search(r'\b'+m+r'\b',announced,re.I)]
                header=r['title'][:1200].lower()
                if len(stated)==1 and stated[0] not in header and any(m in header for m in months):
                    r.update(state='Inconclusivo — possível divergência de mês',action='Conferir mês do demonstrativo; não alterar pelo nome do arquivo.',status='Pendente')
        if not public and r['status']=='Concluído':
            if DOC.search(base):r.update(state='Disponível — vínculo público a conferir',action='Conferir identificação e necessidade de vínculo no Portal.',status='Pendente')
            else:r.update(state='Referência interna disponível',action='',status='Concluído')
        if not public and re.search(r'sint[eé]tic',unquote(base),re.I):
            r.update(state='Histórico — artefato de ensaio fora do Portal atual',action='Preservar referência histórica; não providenciar publicação de ensaio retirado.',status='Concluído')
        mod=old.get('Módulo') or item['module']
        mod={'Convenios':'Convênios','Licitacoes':'Licitações','RecursosHumanos':'RH','PlanejamentoPrestacaoContas':'Planejamento','LGPDeGovernoDigital':'LGPD','Diarias':'Diárias'}.get(mod,mod)
        if mod=='Controle interno':mod=item['module']
        pending=r['status']!='Concluído'
        if pending:counts[mod]+=1
        if r['state'].startswith('Inconclusivo'):inconclusive+=1
        notes='Contextos de controle preservados na aba Ocorrências. '+('Destino somente em controle/fonte; sem vínculo público localizado.' if not public else '')
        history=old.get('Histórico de conferências') or ''
        if old.get('Situação identificada'):history+=f"\n{old.get('Data da conferência')}: {old.get('Situação identificada')}; status anterior: {old.get('Status')}; evidência: {old.get('Evidência da verificação')}"
        d=dict(old);d.update({'Módulo':mod,'Página ou item':old.get('Página ou item') or item['page'],'Descrição do documento':old.get('Descrição do documento') or item['context'][:1000], 'URL atual':url,'Situação identificada':r['state'],'Providência necessária':old.get('Providência necessária') or r['action'],'Status':('Em andamento' if old.get('Status')=='Em andamento' and pending else r['status']),'Observações':old.get('Observações') or notes,'Nome exato esperado do arquivo':old.get('Nome exato esperado do arquivo') or (unquote(Path(urlsplit(base).path).name) if DOC.search(base) else ''),'URL final após redirecionamento':r['final'],'Evidência da verificação':r['evidence'],'Data da conferência':r['date'],'Histórico de conferências':history.strip()})
        ws.append([d.get(k) for k in heads])
    occ=wb.create_sheet('Ocorrências');occ.append(['URL','Módulo','Página ou fonte','Contexto original / providência / responsável','Origem'])
    seen=set()
    for o in occurrences:
        key=(o['url'],o['page'],o['context'],o['kind'])
        if key in seen:continue
        seen.add(key);occ.append([o['url'],o['module'],o['page'],o['context'][:32000],o['kind']])
    summary=wb.create_sheet('Resumo');summary.append(['Indicador','Resultado'])
    for row in [('Conferência',now),('HTMLs examinados',len(pages)),('URLs únicas verificadas (sem fragmentos)',len(results)),('Ocorrências',len(seen)),('Verificações inconclusivas (linhas)',inconclusive),('Navegador',browser),('Base local resolvida',BASE+'Portal-da-Transparencia/'),('Limitações','GET limitado a 24 MiB; primeiros cinco fólios para identificação de PDF; sem OCR ou análise jurídica. Identificação sem correspondência confirmada é inconclusiva. Recursos do pacote conferidos localmente, sem presumir publicação. Exportações Blob verificadas no recorte inicial quando disponíveis; fluxos autenticados não acessados. Não é conciliação financeira de todos os recortes. Resultados reutilizados do cache conservam sua data; use --reverificar para renovar.'),('Controle','Campos manuais e abas antigas preservados; fontes históricas em Ocorrências. Nenhuma URL do Portal foi alterada.')]:summary.append(row)
    for mod,count in sorted(counts.items()):summary.append(['Pendências — '+mod,count])
    for p in pages:summary.append(['Página examinada',p.relative_to(DIST).as_posix()])
    for sheet in [ws,occ,summary]:
        sheet.freeze_panes='A2';sheet.auto_filter.ref=sheet.dimensions
        for cell in sheet[1]:cell.font=Font(bold=True,color='FFFFFF');cell.fill=PatternFill('solid',fgColor='245C43')
        for row in sheet.iter_rows(min_row=2):
            for c in row:
                if isinstance(c.value,str) and c.value.startswith(('http://','https://')) and '\n' not in c.value:
                    c.hyperlink=c.value;c.font=Font(color='0563C1',underline='single')
                c.alignment=Alignment(vertical='top',wrap_text=True)
                if isinstance(c.value,str) and c.value.startswith('='):c.data_type='s'
                if isinstance(c.value,str) and c.value in ['Pendente','Em andamento','Concluído']:c.fill=PatternFill('solid',fgColor={'Pendente':'FCE4D6','Em andamento':'FFF2CC','Concluído':'E2EFDA'}[c.value])
            sheet.row_dimensions[row[0].row].height=58
        for i in range(1,sheet.max_column+1):sheet.column_dimensions[get_column_letter(i)].width=24 if i==1 else 55
    path.parent.mkdir(exist_ok=True,parents=True);wb.save(path)
    print(json.dumps({'xlsx':str(path),'paginas':len(pages),'urls':len(results),'inconclusivas':inconclusive,'pendencias_por_modulo':dict(counts),'navegador':browser},ensure_ascii=False))

def main():
    parser=argparse.ArgumentParser(description=__doc__);parser.add_argument('--reverificar',action='store_true');parser.add_argument('--arquivo',type=Path);args=parser.parse_args()
    WORK.mkdir(parents=True,exist_ok=True)
    candidates=[p for p in ROOT.rglob('controle-publicacao-pdfs-petp-2026.xlsx') if not any(x in p.parts for x in ['.venv','trabalho','.git'])]
    path=args.arquivo or (candidates[0] if candidates else ROOT/'controle/controle-publicacao-documentos.xlsx')
    # Index é somente lido para comprovar preservação.
    index=hashlib.sha256((ROOT/'.git/index').read_bytes()).hexdigest()
    protected={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in DIST.rglob('*') if p.is_file()}
    pages,occ=collect()
    signature=hashlib.sha256(('exportacoes-v3'+''.join(protected.values())).encode()).hexdigest();domcache=WORK/'dom-inicial.json'
    previous=json.loads(domcache.read_text(encoding='utf-8')) if domcache.exists() else {}
    if previous.get('signature')==signature and not args.reverificar:
        dynamic,browser=previous['occurrences'],previous['browser']
    else:
        dynamic,browser=rendered(pages)
        domcache.write_text(json.dumps(dict(signature=signature,occurrences=dynamic,browser=browser),ensure_ascii=False),encoding='utf-8')
    occ+=dynamic
    urls=sorted({urldefrag(o['url'])[0] for o in occ if o['url']});results={}
    print(f'{len(pages)} HTMLs; {len(urls)} destinos únicos. {browser}',flush=True)
    with concurrent.futures.ThreadPoolExecutor(max_workers=4) as pool:
        futures={pool.submit(check,u,args.reverificar):u for u in urls}
        for f in concurrent.futures.as_completed(futures):
            u=futures[f];results[u]=f.result()
            if len(results)%50==0:print(f'Conferidos {len(results)}/{len(urls)} destinos.',flush=True)
    write_book(path,occ,results,pages,browser)
    assert index==hashlib.sha256((ROOT/'.git/index').read_bytes()).hexdigest(),'Índice alterado externamente durante a tarefa.'
    assert all(Path(p).exists() and hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in protected.items()),'Conteúdo público alterado externamente durante a tarefa.'

if __name__=='__main__':main()
