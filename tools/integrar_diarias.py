"""Gera a consulta existente de diárias e seu CSV; --check confere sem gravar."""
import argparse
import csv
import hashlib
import io
import json
import re
from collections import Counter, defaultdict
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from html import escape
from pathlib import Path

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'dados-entrada/diarias/DIARIAS_CONSOLIDADAS_2023_2026.xlsx'
PAGE = ROOT / 'dist/Portal-da-Transparencia/Diarias/Relacao/index.html'
CSV_PATH = ROOT / 'dist/dados/diarias/relacao-diarias.csv'
HEADERS = ['ANO', 'MÊS', 'NOME DO SERVIDOR', 'CARGO', 'DESTINO', 'DESCRIÇÃO DO EVENTO',
           'DATA DE SAÍDA', 'DATA DE RETORNO', 'Nº DE DIÁRIAS', 'TOTAL RECEBIDO', 'SITUAÇÃO']
MONTHS = ['JANEIRO', 'FEVEREIRO', 'MARÇO', 'ABRIL', 'MAIO', 'JUNHO', 'JULHO',
          'AGOSTO', 'SETEMBRO', 'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO']
ITEMS = [None, None, 'nomeBeneficiario', 'cargoBeneficiario', 'localDestino',
         'motivoAfastamento', 'dataInicioAfastamento', 'dataFimAfastamento', 'numeroDiaria', 'valorTotal', None]


def number(value):
    if value is None or isinstance(value, bool):
        return None
    text = str(value).strip().replace('R$', '').replace('\xa0', '').replace(' ', '')
    if ',' in text:
        text = text.replace('.', '').replace(',', '.')
    try:
        result = Decimal(text)
        return result if result.is_finite() else None
    except InvalidOperation:
        return None


def numeric_text(value, money=False):
    places = max(2 if money else 0, -value.as_tuple().exponent)
    return format(value, f'.{places}f').replace('.', ',')


def monetary(value):
    text = numeric_text(value, True)
    whole, fraction = text.split(',')
    sign = '-' if whole.startswith('-') else ''
    whole = format(int(whole.lstrip('-')), ',').replace(',', '.')
    return f'R$ {sign}{whole},{fraction}'


def date_text(value):
    if isinstance(value, (date, datetime)):
        return value.strftime('%d/%m/%Y'), True
    text = '' if value is None else str(value)
    for fmt in ('%d/%m/%Y', '%Y-%m-%d'):
        try:
            return datetime.strptime(text, fmt).strftime('%d/%m/%Y'), True
        except ValueError:
            pass
    return text, False


def read_source():
    book = load_workbook(SOURCE, read_only=True, data_only=False)
    result, pending, sheets = [], [], []
    try:
        for sheet in book:
            values = sheet.iter_rows(values_only=True)
            header = list(next(values))
            if header != HEADERS:
                raise ValueError(f'Cabeçalho inesperado em {sheet.title}; origem não alterada.')
            count = 0
            for row_number, raw in enumerate(values, 2):
                if all(v is None for v in raw):
                    continue
                if any(isinstance(v, str) and v.startswith('=') for v in raw):
                    raise ValueError(f'Fórmula não avaliada: {sheet.title}!{row_number}.')
                if str(raw[0]) not in ('2023', '2024', '2025', '2026') or raw[1] not in MONTHS:
                    raise ValueError(f'Referência não reconhecida: {sheet.title}!{row_number}.')
                out = ['' if v is None else str(v) for v in raw]
                valid = [v is not None and str(v).strip() != '' for v in raw]
                for index in (6, 7):
                    out[index], valid[index] = date_text(raw[index])
                    if not valid[index]:
                        pending.append(dict(aba=sheet.title, linha=row_number, coluna=HEADERS[index],
                                            valor=out[index], motivo='Data inválida ou ausente; preservada sem item Serigy de data.'))
                decimals = {}
                for index in (8, 9):
                    value = number(raw[index])
                    decimals[index] = value
                    valid[index] = value is not None
                    if value is None:
                        pending.append(dict(aba=sheet.title, linha=row_number, coluna=HEADERS[index],
                                            valor=out[index], motivo='Número não interpretado; valor literal preservado.'))
                    else:
                        out[index] = numeric_text(value, index == 9)
                        if index == 9 and value.as_tuple().exponent < -2:
                            pending.append(dict(aba=sheet.title, linha=row_number, coluna=HEADERS[index],
                                                valor=str(raw[index]), motivo='Mais de duas casas decimais; precisão preservada.'))
                for index in range(11):
                    if raw[index] is None or str(raw[index]).strip() == '':
                        pending.append(dict(aba=sheet.title, linha=row_number, coluna=HEADERS[index],
                                            motivo='Campo ausente; sem item Serigy de dado.'))
                result.append(dict(values=out, valid=valid, money=decimals[9],
                                   year=out[0], month=f'{MONTHS.index(raw[1])+1:02d}',
                                   sheet=sheet.title, source_row=row_number))
                count += 1
            sheets.append(dict(aba=sheet.title, registros=count, colunas=len(header)))
    finally:
        book.close()
    if not result:
        raise ValueError('Origem sem registros; página não alterada.')
    return result, pending, sheets


def csv_text(rows):
    stream = io.StringIO(newline='')
    writer = csv.writer(stream, delimiter=';', quoting=csv.QUOTE_ALL)
    writer.writerow(HEADERS)
    writer.writerows(r['values'] for r in rows)
    return '\ufeff' + stream.getvalue()


def main_html(rows):
    years = sorted({r['year'] for r in rows}, reverse=True)
    year_options = ''.join(f'<option value="{y}" serigyValue="{y}">{y}</option>' for y in years)
    month_options = ''.join(f'<option value="{i:02d}">{name.title()}</option>' for i, name in enumerate(MONTHS, 1))
    rendered = []
    for row in rows:
        cells = []
        for index, text in enumerate(row['values']):
            item = f' serigyItem="{ITEMS[index]}"' if ITEMS[index] and row['valid'][index] else ''
            label = escape(text).replace('\r', '&#13;').replace('\n', '&#10;') if text else 'Não informado'
            if index == 2:
                label = '<span data-search-text>' + label + '</span>'
            if index == 9 and row['money'] is not None:
                label = monetary(row['money'])
            css = ' class="money"' if index == 9 else ''
            cells.append(f'<td{item}{css}>{label}</td>')
        export = escape(json.dumps(row['values'], ensure_ascii=False), quote=True)
        rendered.append(f'<tr data-year="{row["year"]}" data-month="{row["month"]}" data-source-row="{row["source_row"]}" data-export-values="{export}">'+''.join(cells)+'</tr>')
    periods = sorted({(r['year'], r['month']) for r in rows})
    period_label = lambda p: MONTHS[int(p[1])-1].title() + '/' + p[0]
    return f'''<main id="conteudo" class="container" serigySchema="https://serigy.tce.se.gov.br/comum/diarias.html">
    <h1>Relação de Diárias</h1>
    <p class="lead">Consulte os beneficiários, os períodos de afastamento, os destinos e os valores das diárias.</p>
    <form id="filtro-relacao-diarias" class="filter-panel" role="search" aria-label="Consultar relação de diárias" data-table-filter data-target="tabela-relacao-diarias" serigyItem="filtroPesquisa">
      <div class="field"><label for="nome-relacao-diarias">Nome do beneficiário</label><input id="nome-relacao-diarias" name="palavra" type="search"></div>
      <div class="field"><label for="mes-relacao-diarias">Mês de referência</label><select id="mes-relacao-diarias" name="mes"><option value="">Todos os meses</option>{month_options}</select></div>
      <div class="field"><label for="ano-relacao-diarias">Ano de referência</label><select id="ano-relacao-diarias" name="ano" serigyItem="historicoExercicio"><option value="">Todos os anos</option>{year_options}</select></div>
      <div class="filter-actions"><button id="pesquisar-relacao-diarias" type="submit" serigyExecuteWaitSection="resultados-relacao-diarias">Pesquisar</button><button type="reset" class="secondary" serigyClearForm>Limpar filtros</button></div>
    </form>
    <section id="resultados-relacao-diarias" aria-labelledby="titulo-relacao-diarias" serigyLoadSection="resultados-relacao-diarias">
      <h2 id="titulo-relacao-diarias">Diárias por afastamento</h2>
      <div class="exports"><a id="exportar-relacao-diarias" href="/dados/diarias/relacao-diarias.csv" download="relacao-diarias.csv" serigyItem="salvarRelatorio" serigyValue="csv">Exportar registros selecionados em CSV</a></div>
      <p class="result-count" data-count-for="tabela-relacao-diarias" aria-live="polite">{len(rows)} registro(s) encontrado(s).</p>
      <p id="diarias-vazio" hidden>Nenhum registro encontrado para os filtros selecionados.</p>
      <noscript><p>Sem JavaScript, a tabela e o CSV apresentam todos os registros.</p><style>#filtro-relacao-diarias {{ display: none; }}</style></noscript>
      <div class="table-wrap" tabindex="0" role="region" aria-label="Relação de diárias com rolagem horizontal">
        <table id="tabela-relacao-diarias">
          <caption>Relação de diárias por afastamento</caption>
          <thead><tr><th scope="col">Ano de referência</th><th scope="col">Mês de referência</th><th scope="col">Nome completo do beneficiário</th><th scope="col">Cargo/função</th><th scope="col">Destino (cidade)</th><th scope="col">Motivo do afastamento</th><th scope="col">Início do afastamento</th><th scope="col">Fim do afastamento</th><th scope="col">Número de diárias</th><th scope="col">Valor total recebido</th><th scope="col">Situação na fonte</th></tr></thead>
          <tbody id="registros-relacao-diarias">
{chr(10).join(rendered)}
          </tbody>
        </table>
      </div>
    </section>
    <div class="update-info" serigyItem="dataUltimaAtualizacao">Informações atualizadas em <time datetime="2026-10-22">22/10/2026</time></div>
  </main>'''


def generate(check=False):
    checksum = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    rows, pending, sheets = read_source()
    page = PAGE.read_text(encoding='utf-8')
    html, count = re.subn(r'<main\b.*?</main>', lambda _: main_html(rows), page, count=1, flags=re.S)
    if count != 1:
        raise ValueError('Região principal não encontrada; página não alterada.')
    html = html.replace('<body class="awaiting-information">', '<body>')
    for asset in ('<link rel="stylesheet" href="/assets/css/diarias.css">',
                  '<script defer src="/assets/js/tabelas.js"></script>',
                  '<script defer src="/assets/js/diarias.js"></script>'):
        if asset not in html:
            html = html.replace('</head>', '  ' + asset + '\n</head>', 1)
    files = {PAGE: html.encode('utf-8'), CSV_PATH: csv_text(rows).encode('utf-8')}
    if check:
        for path, content in files.items():
            if not path.exists() or path.read_bytes() != content:
                raise ValueError('Arquivo divergente da origem: ' + str(path.relative_to(ROOT)))
    else:
        for path, content in files.items():
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content)
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == checksum
    amounts = defaultdict(Decimal)
    for row in rows:
        if row['money'] is not None:
            amounts[row['year']] += row['money']
    report = dict(origem_sha256=checksum, abas=sheets, registros=len(rows),
                  por_ano=dict(Counter(r['year'] for r in rows)),
                  totais_por_ano={y: str(v) for y, v in sorted(amounts.items())}, pendencias=pending)
    if not check:
        local = ROOT / 'trabalho/diarias-conferencia.json'
        local.parent.mkdir(exist_ok=True)
        local.write_text(json.dumps(report, ensure_ascii=False, indent=2)+'\n', encoding='utf-8')
    print(json.dumps(report, ensure_ascii=False, indent=2))
    return report


if __name__ == '__main__':
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    generate(parser.parse_args().check)
