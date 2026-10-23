"""Concilia todos os campos do HTML e CSV com a planilha original de diárias."""
import csv
import hashlib
import re
from collections import defaultdict
from datetime import datetime
from decimal import Decimal
from html.parser import HTMLParser

from openpyxl import load_workbook
from integrar_diarias import ROOT, SOURCE, PAGE, CSV_PATH, generate


class Table(HTMLParser):
    def __init__(self, text):
        super().__init__()
        self.rows, self.row, self.cell = [], None, None
        self.feed(text)

    def handle_starttag(self, tag, attrs):
        attrs = dict(attrs)
        if tag == 'tr' and 'data-source-row' in attrs:
            self.row = dict(attrs=attrs, cells=[])
        elif tag == 'td' and self.row is not None:
            self.cell = dict(attrs=attrs, text='')

    def handle_data(self, data):
        if self.cell is not None:
            self.cell['text'] += data

    def handle_endtag(self, tag):
        if tag == 'td' and self.cell is not None:
            self.row['cells'].append(self.cell)
            self.cell = None
        elif tag == 'tr' and self.row is not None:
            self.rows.append(self.row)
            self.row = None


def decimal(value):
    if isinstance(value, (int, float)):
        return Decimal(str(value))
    text = re.sub(r'R\$|\s', '', str(value))
    return Decimal(text.replace('.', '').replace(',', '.')) if ',' in text else Decimal(text)


def main():
    checksum = hashlib.sha256(SOURCE.read_bytes()).hexdigest()
    generate(check=True)
    book = load_workbook(SOURCE, read_only=True, data_only=False)
    assert book.sheetnames == ['Plan1']
    raw = list(book['Plan1'].iter_rows(values_only=True))
    header, source = list(raw[0]), raw[1:]
    book.close()
    with CSV_PATH.open(encoding='utf-8-sig', newline='') as file:
        exported = list(csv.reader(file, delimiter=';'))
    assert exported.pop(0) == header
    text = PAGE.read_text(encoding='utf-8')
    rows = Table(text).rows
    assert len(source) == len(exported) == len(rows) == 208
    assert text.count('<table ') == 1 and text.count('serigyItem="dataUltimaAtualizacao"') == 1 and '22/10/2026' in text
    assert text.index('src="/assets/js/tabelas.js"') < text.index('src="/assets/js/diarias.js"')
    source_totals, csv_totals, html_totals = [defaultdict(Decimal) for _ in range(3)]
    start_dates = 0
    for line, (record, out, row) in enumerate(zip(source, exported, rows), 2):
        cells = row['cells']
        assert len(out) == len(cells) == 11 and 'hidden' not in row['attrs']
        assert row['attrs']['data-source-row'] == str(line)
        assert row['attrs']['data-year'] == str(record[0])
        for index in (0, 1, 2, 3, 4, 5, 10):
            assert str(record[index]) == out[index] == cells[index]['text'], (line, index)
        for index in (6, 7):
            expected = record[index].strftime('%d/%m/%Y') if isinstance(record[index], datetime) else record[index]
            assert out[index] == cells[index]['text'] == expected
        for index in (8, 9):
            assert decimal(record[index]) == decimal(out[index]) == decimal(cells[index]['text']), (line, index)
        for index, item in ((2,'nomeBeneficiario'), (3,'cargoBeneficiario'), (4,'localDestino'),
                            (5,'motivoAfastamento'), (7,'dataFimAfastamento'), (8,'numeroDiaria'), (9,'valorTotal')):
            assert cells[index]['attrs'].get('serigyitem') == item
        if line == 13:
            assert cells[6]['text'] == '30/03/20213' and 'serigyitem' not in cells[6]['attrs']
        else:
            assert cells[6]['attrs'].get('serigyitem') == 'dataInicioAfastamento'
            start_dates += 1
        source_totals[str(record[0])] += decimal(record[9])
        csv_totals[out[0]] += decimal(out[9])
        html_totals[cells[0]['text']] += decimal(cells[9]['text'])
    assert start_dates == 207
    assert source_totals == csv_totals == html_totals
    assert dict(source_totals) == {'2023':Decimal('41025.003'), '2024':Decimal('35450'),
                                   '2025':Decimal('75925'), '2026':Decimal('55300')}
    assert sum(source_totals.values()) == Decimal('207700.003')
    assert exported[5][9] == '25,003' and rows[5]['cells'][9]['text'] == 'R$ 25,003'
    assert len(set(map(tuple, exported))) == 208
    assert hashlib.sha256(SOURCE.read_bytes()).hexdigest() == checksum
    print('PASS: 208 linhas/11 campos conciliados no XLSX, HTML e CSV; totais exatos por ano e total 207700.003; precisão de J7 e literal de G13 preservados; 207 datas iniciais válidas; XLSX intacto.')


if __name__ == '__main__':
    main()
