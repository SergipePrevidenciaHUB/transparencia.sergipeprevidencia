"""Integra a relação de licitações sem inferir abertura ou homologação."""
from pathlib import Path
import datetime as dt
from decimal import Decimal, InvalidOperation
import html,json,re,sys
from openpyxl import load_workbook
ROOT=Path(__file__).resolve().parents[1]
PAGE=ROOT/'dist/Portal-da-Transparencia/Licitacoes/index.html'
SOURCE=ROOT/'dados-entrada/licitacoes/relacao-licitacoes.xlsx'
TEST_START='<!-- INICIO ENSAIO LICITACAO 8.1.6 -->'
TEST_END='<!-- FIM ENSAIO LICITACAO 8.1.6 -->'

def esc(v):return html.escape(str(v),quote=True).replace('\r','&#13;').replace('\n','&#10;')
def date(v):
    if isinstance(v,(dt.datetime,dt.date)):return v.strftime('%d/%m/%Y')
    for fmt in ['%d/%m/%Y','%Y-%m-%d','%Y-%m-%d %H:%M:%S']:
        try:return dt.datetime.strptime(str(v),fmt).strftime('%d/%m/%Y')
        except ValueError:pass
    return None
def money(v):
    if v is None or str(v).strip() in ('','-'):return None
    try:
        d=Decimal(str(v));assert d.is_finite()
        return format(d,'f').replace('.',',')
    except (InvalidOperation,AssertionError):return None
def records():
    w=load_workbook(SOURCE,data_only=False);result=[]
    for s in w:
        for cells in s.iter_rows(min_row=2):
            if not any(c.value is not None for c in cells):continue
            vals=[c.value for c in cells]; links=[]
            for c in cells:
                if c.hyperlink and c.hyperlink.target:links.append({'coluna':c.column_letter,'texto':c.value,'url':c.hyperlink.target})
                elif isinstance(c.value,str) and re.fullmatch(r'https?://\S+',c.value.strip()):links.append({'coluna':c.column_letter,'texto':c.value,'url':c.value.strip()})
            result.append({'aba':s.title,'linha':cells[0].row,'valores':vals,'links':links})
    return result
def atas_section():
    source=json.loads((ROOT/'dados-entrada/contratos/contratos-extraidos.json').read_text(encoding='utf-8'))
    entries=[]
    for r in source['atas_para_revisao']:
        # Natureza explícita no campo modalidade/termo ou no objeto, nunca no nome do PDF.
        evidence=str(r['celulas_origem'].get('B') or '')+' '+str(r['objeto_literal'] or '')
        if not r['url_integra'] or 'adesão' not in evidence.lower():continue
        entries.append('<li data-ata-origem="'+str(r['linha_origem'])+'"><p>'+esc(r['identificacao_literal'])+' — '+esc(r['celulas_origem'].get('B') or '')+'</p><p>'+esc(r['contratado_literal'])+' — '+esc(r['objeto_literal'])+'</p><a serigyItem="atasLicitacao" href="'+esc(r['url_integra'])+'">Documento de adesão/SRP</a></li>')
    return '<section id="atas-adesao-srp" aria-labelledby="titulo-atas-adesao"><h2 id="titulo-atas-adesao">Adesões a atas de registro de preços</h2><ul>'+''.join(entries)+'</ul></section>'

def is_pregao(v):
    return 'PREG' in str(v[9]).upper()

def sequence(v):
    match=re.fullmatch(r'[^0-9]*(\d+)\s*/\s*\d{4}',str(v[6]).strip())
    if not match:raise ValueError('Identificador sem sequência inequívoca: '+str(v[6]))
    return int(match[1])

def consultation_dates(v):
    return (date(v[14]),None) if is_pregao(v) else (None,date(v[15]))

def render(synthetic=None):
    old=PAGE.read_text(encoding='utf-8')
    if synthetic:raise ValueError("Ensaio sintético de Licitações descontinuado; usar registros reais.")
    data=sorted(records(),key=lambda r:(str(r['valores'][0]),str(r['valores'][9]),sequence(r['valores'])));rows=[];years=set();modal=set();status=set();openings=0
    docitems={'AD':'editalLicitacao','AE':'docFasesLicitacao','AG':'docDispensaLicitacao'}
    doclabels={'AD':'Edital','AE':'Documentos das fases de licitação','AG':'Documentos de dispensa/inexigibilidade'}
    for r in data:
        v=r['valores'];year=str(v[0]) if v[0] is not None else '';m=str(v[9] or '');s=str(v[5] or '')
        if year:years.add(year)
        if m:modal.add(m)
        if s:status.add(s)
        opening,start=consultation_dates(v);reference=opening or start;openings+=bool(opening)
        docs=[]
        for l in r['links']:
            item=docitems.get(l['coluna']);attr=' serigyItem="'+item+'"' if item else ''
            label=doclabels.get(l['coluna'],'Link da fonte ('+l['coluna']+')')
            docs.append('<li><a'+attr+' href="'+esc(l['url'])+'">'+esc(label)+'</a></li>')
        def cell(value,item=None):
            if value is None or str(value).strip() in ('','-'):return '<td></td>'
            return '<td'+(' serigyItem="'+item+'"' if item else '')+'>'+esc(value)+'</td>'
        rows.append('<tr data-source-row="'+str(r['linha'])+'" data-year="'+esc(year)+'" data-modalidade="'+esc(m)+'" data-situacao="'+esc(s)+'" data-mes="'+(reference[3:5] if reference else '')+'">'+cell(year)+cell(v[6])+cell(sequence(v),'sequencia')+cell(v[7])+cell(v[9],'modalidadeLicitacao')+cell(v[11],'objetoLicitacao')+cell(opening,'dataLicitacao')+cell(start)+cell(money(v[23]),'valorLicitacao' if money(v[23]) is not None else None)+cell(money(v[24]),'valorLicitacao' if money(v[23]) is None else None)+cell(v[5],'faseLicitacao')+'<td>'+('<ul>'+''.join(docs)+'</ul>' if docs else '')+'</td></tr>')
    def select(name,label,values,hist=False):
        return '<div class="field"><label for="'+name+'-licitacoes">'+label+'</label><select name="'+name+'" id="'+name+'-licitacoes"'+(' serigyItem="historicoExercicio" serigyExecuteWaitSection="licitacoesTabela"' if hist else '')+'><option value="">Todos</option>'+''.join('<option value="'+esc(v)+'"'+(' serigyValue="'+esc(v)+'"' if hist else '')+'>'+esc(v)+'</option>' for v in sorted(values,reverse=hist))+'</select></div>'
    main='<main id="conteudo" class="container" serigySchema="https://serigy.tce.se.gov.br/comum/licitacoes.html"><h1>Licitações</h1><p>Relação de procedimentos e documentos. O ano é o informado na fonte. Valores de referência e arrematados são apresentados separadamente, em reais.</p><form id="filtro-licitacoes" class="filter-panel" serigyItem="filtroPesquisa" aria-label="Consultar licitações">'+select('ano','Ano',years,True)+select('modalidade','Modalidade',modal)+select('situacao','Situação',status)+select('mes','Mês',{d[3:5] for r in data for d in [next((x for x in consultation_dates(r['valores']) if x),None)] if d})+'<div class="field"><label for="texto-licitacoes">Pesquisa textual</label><input type="search" name="texto" id="texto-licitacoes"></div><button type="submit" serigyExecuteWaitSection="licitacoesTabela">Pesquisar</button><button type="reset" serigyClearForm>Limpar filtros</button></form><section serigyLoadSection="licitacoesTabela" aria-label="Resultados"><p id="contagem-licitacoes" role="status" aria-live="polite">'+str(len(data))+' procedimentos.</p><button id="exportar-licitacoes" type="button" serigyItem="salvarRelatorio" serigyValue="csv">CSV do recorte</button><div class="table-wrap"><table id="tabela-licitacoes"><caption>Procedimentos e documentos vinculados</caption><thead><tr>'+''.join('<th scope="col">'+h+'</th>' for h in ['Ano','Número identificador','Sequência','Processo','Modalidade','Objeto','Data de abertura','Data de início','Valor de referência (R$)','Valor arrematado (R$)','Situação / Etapa','Documentos'])+'</tr></thead><tbody>'+'\n'.join(rows)+'</tbody></table></div></section><p class="update-info">Informações atualizadas em <span serigyItem="dataUltimaAtualizacao">22/10/2026</span></p></main>'
    main=main.replace('</main>',atas_section()+'</main>')
    new=re.sub(r'<main\b.*?</main>',lambda _:main,old,count=1,flags=re.S)
    if '/assets/js/licitacoes.js' not in new:new=new.replace('</head>','<script defer src="/assets/js/licitacoes.js"></script>\n</head>')
    return new,data,openings
if __name__=='__main__':
    if '--ativar-teste' in sys.argv and '--desativar-teste' in sys.argv:raise SystemExit('Escolha apenas ativar ou desativar o ensaio.')
    synthetic=True if '--ativar-teste' in sys.argv else False if '--desativar-teste' in sys.argv else None
    content,data,openings=render(synthetic)
    if '--check' in sys.argv:assert PAGE.read_text(encoding='utf-8')==content
    else:PAGE.write_text(content,encoding='utf-8')
    from collections import Counter
    print(json.dumps({'registros':len(data),'anos':dict(Counter(str(r['valores'][0]) for r in data)),'aberturas_validas':openings,'links':dict(Counter(l['coluna'] for r in data for l in r['links']))},ensure_ascii=True))
