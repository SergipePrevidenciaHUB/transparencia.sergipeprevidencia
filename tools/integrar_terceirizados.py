"""Publica somente as cinco colunas autorizadas da relação mensal de terceirizados."""
from pathlib import Path
from collections import Counter
from html import escape
import argparse
import csv
import io
import json
import re

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
SOURCE = ROOT / 'dados-entrada/rh/terceirizados.xlsx'
PAGE = ROOT / 'dist/Portal-da-Transparencia/RecursosHumanos/Terceirizados/index.html'
CSV = ROOT / 'dist/dados/recursos-humanos/terceirizados.csv'
HEADERS = ['Ano', 'Mês', 'Nome completo', 'Função ou atividade exercida', 'Empresa empregadora']
MONTHS = ['JANEIRO', 'FEVEREIRO', 'MARÇO', 'ABRIL', 'MAIO', 'JUNHO', 'JULHO', 'AGOSTO', 'SETEMBRO', 'OUTUBRO', 'NOVEMBRO', 'DEZEMBRO']
SCHEMA = 'https://serigy.tce.se.gov.br/comum/recursos-humanos.html'


def load():
    book = load_workbook(SOURCE, read_only=True, data_only=False)
    try:
        sheet = book['Plan1']
        assert tuple(next(sheet.values))[:5] == ('ANO', 'MÊS', 'NOME', 'CARGO', 'EMPRESA')
        records = []
        for number, row in enumerate(sheet.iter_rows(min_row=2, values_only=True), 2):
            values = [str(int(row[0]))] + [str(v) if v is not None else '' for v in row[1:5]]
            assert values[1] in MONTHS, f'Mês não reconhecido na linha {number}'
            records.append((number, values))
        return records
    finally:
        book.close()


def csv_text(records):
    output = io.StringIO(newline='')
    writer = csv.writer(output, delimiter=';', quoting=csv.QUOTE_ALL, lineterminator='\r\n')
    writer.writerow(HEADERS)
    writer.writerows(values for _, values in records)
    return '\ufeff' + output.getvalue()


def render(records):
    years = ''.join(f'<option value="{y}" serigyValue="{y}">{y}</option>' for y in sorted({v[0] for _, v in records}, reverse=True))
    months = ''.join(f'<option value="{m}">{m.title()}</option>' for m in MONTHS if any(v[1] == m for _, v in records))
    rows = []
    for number, values in records:
        attrs = escape(json.dumps(values, ensure_ascii=False), quote=True)
        cells = []
        for value, item in zip(values, ['', '', 'nomeTerceirizado', 'funcaoTerceirizado', 'empresaTerceirizado']):
            present = bool(value.strip())
            marker = f' serigyItem="{item}"' if item and present else ''
            cells.append(f'<td{marker}>{escape(value) if present else "Não informado"}</td>')
        rows.append(f'          <tr data-source-row="{number}" data-year="{values[0]}" data-month="{values[1]}" data-export-values="{attrs}">' + ''.join(cells) + '</tr>')
    headings = ''.join(f'<th scope="col">{h}</th>' for h in HEADERS)
    return f'''<main id="conteudo" class="container" serigySchema="{SCHEMA}">
    <h1>Terceirizados</h1>
    <p class="lead">Relação mensal de trabalhadores terceirizados, suas atividades e empresas empregadoras.</p>
    <form id="filtro-terceirizados" class="filter-panel" role="search" aria-label="Consultar terceirizados" serigyItem="filtroPesquisa">
      <div class="field"><label for="ano-terceirizados">Ano da relação</label><select id="ano-terceirizados" name="ano" serigyItem="historicoExercicio"><option value="">Todos</option>{years}</select></div>
      <div class="field"><label for="mes-terceirizados">Mês da relação</label><select id="mes-terceirizados" name="mes"><option value="">Todos</option>{months}</select></div>
      <div class="filter-actions"><button id="pesquisar-terceirizados" type="submit" serigyExecuteWaitSection="tabela-terceirizados">Pesquisar</button><button type="reset" class="secondary" serigyClearForm>Limpar filtros</button></div>
    </form>
    <section id="resultados-terceirizados" aria-labelledby="titulo-terceirizados">
      <h2 id="titulo-terceirizados">Relação de terceirizados</h2>
      <div class="exports"><strong>Baixar registros selecionados:</strong> <a id="exportar-terceirizados" href="/dados/recursos-humanos/terceirizados.csv" download serigyItem="salvarRelatorio" serigyValue="csv">CSV</a></div>
      <p class="result-count" data-count-for="tabela-terceirizados" aria-live="polite">{len(records)} registro(s) encontrado(s).</p>
      <p id="terceirizados-vazio" hidden>Nenhum registro encontrado para os filtros selecionados.</p>
      <div class="table-wrap"><table id="tabela-terceirizados" serigyLoadSection="tabela-terceirizados">
        <caption>Trabalhadores terceirizados por período da relação</caption>
        <thead><tr>{headings}</tr></thead>
        <tbody id="registros-terceirizados">
{chr(10).join(rows)}
        </tbody>
      </table></div>
    </section>
    <div class="update-info" serigySchema="https://serigy.tce.se.gov.br/comum/recursos-humanos.html" serigyItem="dataUltimaAtualizacao">Informações atualizadas em <time datetime="2026-10-22">22/10/2026</time></div>
  </main>'''


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument('--check', action='store_true')
    args = parser.parse_args()
    records = load()
    page = PAGE.read_text(encoding='utf-8')
    page = re.sub(r'<main id="conteudo".*?</main>', lambda _: render(records), page, flags=re.S)
    page = page.replace('<body class="awaiting-information">', '<body>')
    if '/assets/js/terceirizados.js' not in page:
        page = page.replace('</head>', '  <script defer src="/assets/js/terceirizados.js"></script>\n</head>')
    for path, content in [(PAGE, page), (CSV, csv_text(records))]:
        if args.check:
            assert path.read_bytes() == content.encode('utf-8'), f'Divergência: {path}'
        else:
            path.parent.mkdir(parents=True, exist_ok=True)
            path.write_bytes(content.encode('utf-8'))
    print(json.dumps({'registros': len(records), 'anos': dict(Counter(v[0] for _, v in records)), 'lacunas': [{'linha': n, 'colunas': [HEADERS[i] for i, value in enumerate(v) if not value.strip()]} for n, v in records if any(not value.strip() for value in v)]}, ensure_ascii=False))


if __name__ == '__main__':
    main()
