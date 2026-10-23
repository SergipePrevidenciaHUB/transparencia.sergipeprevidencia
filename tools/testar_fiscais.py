"""Concilia associações com XLSX e testa filtros/CSV no Chrome local."""
from pathlib import Path
from openpyxl import load_workbook
from http.server import ThreadingHTTPServer,SimpleHTTPRequestHandler
from functools import partial
from threading import Thread
from urllib.request import urlopen
import re,json,subprocess,tempfile,html,csv,io,hashlib
ROOT=Path(__file__).resolve().parents[1];W=ROOT/'trabalho/fiscais';W.mkdir(parents=True,exist_ok=True);protected={str(p):hashlib.sha256(p.read_bytes()).hexdigest() for p in (ROOT/'dist').rglob('*') if p.is_file()};run=Path(tempfile.mkdtemp(dir=W));data=[];roles=['Fiscal','Fiscal substituto','Gestor','Gestor substituto']
for i,row in enumerate(load_workbook(ROOT/'dados-entrada/contratos/fiscais-e-gestores.xlsx',data_only=True).active.iter_rows(values_only=True),1):
 if i in [1,107] or not row[1]:continue
 inst=re.sub(r'\s*\*+\s*$','',re.sub(r'^\s*\*+\s*','',str(row[1]).strip()));ys=sorted(set(re.findall(r'\b(?:19|20)\d{2}\b',inst)))
 for role,n in zip(roles,row[7:11]):
  if n is not None and str(n).strip() not in ['', '-']:data.append([ys[-1] if ys else 'Não identificado',inst,str(row[2] or '').strip() or 'Não informado',role,str(n).strip()])
script=r"""<script>window.addEventListener('load',async()=>{const expected=SPEC,failures=[];let checks=0;function check(v,m){checks++;if(!v)failures.push(m);}const f=document.getElementById('filtro-fiscais'),t=document.getElementById('tabela-fiscais'),rows=Array.from(t.tBodies[0].rows);check(rows.length===expected.length,'total');rows.forEach((r,i)=>{check(JSON.stringify(Array.from(r.cells).map(c=>c.textContent.trim()))===JSON.stringify(expected[i]),'associação '+i);check(r.cells[1].getAttribute('serigyItem')==='contratoFiscal'&&r.cells[4].getAttribute('serigyItem')==='fiscalGestor','marcas '+i);});function apply(){f.dispatchEvent(new Event('submit',{cancelable:true,bubbles:true}));}for(const role of ['Fiscal','Fiscal substituto','Gestor','Gestor substituto']){f.elements.funcao.value=role;apply();check(rows.filter(r=>!r.hidden).length===expected.filter(r=>r[3]===role).length,'função '+role);}f.elements.funcao.value='';for(const o of Array.from(f.elements.ano.options)){f.elements.ano.value=o.value;f.dispatchEvent(new Event('change',{bubbles:true}));check(rows.filter(r=>!r.hidden).length===expected.filter(r=>!o.value||r[0]===o.value).length,'ano '+o.value);}const first=expected[0];f.elements.ano.value=first[0];f.elements.nome.value=first[4].toLowerCase();f.elements.funcao.value=first[3];apply();check(rows.filter(r=>!r.hidden).length===expected.filter(r=>r[0]===first[0]&&r[3]===first[3]&&r[4].toLowerCase().includes(first[4].toLowerCase())).length,'combinados');async function capture(){let blob;URL.createObjectURL=b=>{blob=b;return 'blob:teste';};URL.revokeObjectURL=()=>{};HTMLAnchorElement.prototype.click=function(){};document.getElementById('csv-fiscais').click();return await blob.text();}const selected=await capture();f.reset();await new Promise(r=>setTimeout(r,40));check(rows.every(r=>!r.hidden),'reset');const all=await capture();const pre=document.createElement('pre');pre.id='teste';pre.textContent=JSON.stringify({checks,failures,selected,all});document.body.append(pre);});</script>""".replace('SPEC',json.dumps(data,ensure_ascii=False))
route='/Portal-da-Transparencia/Contratos/GestoresFiscais/';page=(ROOT/'dist'/route.strip('/')/'index.html').read_text(encoding='utf-8')
class Handler(SimpleHTTPRequestHandler):
 def log_message(self,*args):pass
 def do_GET(self):
  if '?teste' not in self.path:return super().do_GET()
  self.send_response(200);self.send_header('Content-Type','text/html; charset=utf-8');self.end_headers();self.wfile.write(page.replace('</body>',script+'</body>').encode())
server=ThreadingHTTPServer(('127.0.0.1',0),partial(Handler,directory=str(ROOT/'dist')));Thread(target=server.serve_forever,daemon=True).start();host='http://127.0.0.1:'+str(server.server_port)
try:
 with urlopen(host+route,timeout=10) as r:assert r.status==200
 b=subprocess.run(['C:/Program Files/Google/Chrome/Application/chrome.exe','--headless','--disable-gpu','--no-sandbox','--no-first-run','--user-data-dir='+str(run/'chrome'),'--virtual-time-budget=4000','--dump-dom',host+route+'?teste=1'],capture_output=True,timeout=45);dom=b.stdout.decode('utf-8',errors='replace');(run/'browser.html').write_text(dom,encoding='utf-8');m=re.search(r'<pre id="teste">(.*?)</pre>',dom,re.S);assert m,'Chrome sem resultado';v=json.loads(html.unescape(m[1]));assert not v['failures'],v
 for key in ['selected','all']:
  text=v.pop(key);parsed=list(csv.reader(io.StringIO(text.lstrip(chr(65279))),delimiter=';'));first=data[0];expected=data if key=='all' else [r for r in data if r[0]==first[0] and r[3]==first[3] and first[4].lower() in r[4].lower()];assert parsed[1:]==expected;(run/(key+'.csv')).write_text(text,encoding='utf-8')
finally:server.shutdown();server.server_close()
assert all(hashlib.sha256(Path(p).read_bytes()).hexdigest()==h for p,h in protected.items())
v.update(associacoes=len(data),CSV_conciliado=True,demais_paginas_preservadas=True);(W/'resultado.json').write_text(json.dumps(v,ensure_ascii=False,indent=2),encoding='utf-8');print(json.dumps(v,ensure_ascii=False))
